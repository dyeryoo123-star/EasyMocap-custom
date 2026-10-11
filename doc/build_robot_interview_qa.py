from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_TAB_ALIGNMENT
from pathlib import Path


OUT = Path("doc/机器人控制系统岗位_RGBD项目面试细节问答.docx")

FONT = "Arial Unicode MS"
MONO = "Menlo"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "172B4D"
MUTED = "5F6B7A"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
PALE_YELLOW = "FFF4CE"
RISK_RED = "9B1C1C"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[idx]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def set_run(run, size=10.5, bold=False, color=INK, font=FONT, italic=False):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def shade_paragraph(paragraph, fill, border=None):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    if border:
        p_bdr = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "6")
        left.set(qn("w:color"), border)
        p_bdr.append(left)
        p_pr.append(p_bdr)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run(run, 9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    run = paragraph.add_run(" 页")
    set_run(run, 9, color=MUTED)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

# compact_reference_guide preset tokens
normal = doc.styles["Normal"]
normal.font.name = FONT
normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

for name, size, color, before, after in (
    ("Heading 1", 16, BLUE, 18, 10),
    ("Heading 2", 13, BLUE, 14, 7),
    ("Heading 3", 12, DARK_BLUE, 10, 5),
):
    st = doc.styles[name]
    st.font.name = FONT
    st._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True
    st.paragraph_format.line_spacing = 1.1

for style_name, base, size, color, bold in (
    ("Question", "Heading 2", 12.5, BLUE, True),
    ("AnswerLabel", "Normal", 10.5, DARK_BLUE, True),
    ("CodeBlock", "Normal", 9.2, INK, False),
    ("Callout", "Normal", 10.3, INK, False),
):
    if style_name not in doc.styles:
        st = doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
    else:
        st = doc.styles[style_name]
    st.base_style = doc.styles[base]
    st.font.name = MONO if style_name == "CodeBlock" else FONT
    st._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.font.size = Pt(size)
    st.font.bold = bold
    st.font.color.rgb = RGBColor.from_string(color)
    if style_name == "Question":
        st.paragraph_format.space_before = Pt(12)
        st.paragraph_format.space_after = Pt(5)
        st.paragraph_format.keep_with_next = True
    elif style_name == "CodeBlock":
        st.paragraph_format.left_indent = Inches(0.18)
        st.paragraph_format.right_indent = Inches(0.12)
        st.paragraph_format.space_before = Pt(4)
        st.paragraph_format.space_after = Pt(7)
        st.paragraph_format.line_spacing = 1.05

header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
run = hp.add_run("机器人控制系统岗位｜项目面试问答")
set_run(run, 9, bold=True, color=MUTED)
footer = section.footer
add_page_number(footer.paragraphs[0])


def add_para(text, bold_prefix=None, color=INK, after=6, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = keep
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run(r1, bold=True, color=DARK_BLUE)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run(r2, color=color)
    else:
        r = p.add_run(text)
        set_run(r, color=color)
    return p


def add_callout(label, text, risk=False):
    p = doc.add_paragraph(style="Callout")
    p.paragraph_format.left_indent = Inches(0.14)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_together = True
    shade_paragraph(p, PALE_YELLOW if risk else LIGHT_BLUE, RISK_RED if risk else BLUE)
    r = p.add_run(label + "：")
    set_run(r, bold=True, color=RISK_RED if risk else DARK_BLUE)
    r = p.add_run(text)
    set_run(r, color=INK)
    return p


def add_code(text):
    p = doc.add_paragraph(style="CodeBlock")
    shade_paragraph(p, LIGHT_GRAY)
    for i, line in enumerate(text.splitlines()):
        if i:
            p.add_run().add_break()
        r = p.add_run(line)
        set_run(r, 9.2, font=MONO, color=INK)
    return p


q_no = 0


def add_qa(question, answer, follow=None, boundary=None, code=None):
    global q_no
    q_no += 1
    p = doc.add_paragraph(style="Question")
    r = p.add_run(f"Q{q_no}. {question}")
    set_run(r, 12.5, bold=True, color=BLUE)
    add_para("回答：" + answer, bold_prefix="回答：")
    if code:
        add_code(code)
    if follow:
        add_para("追问时补充：" + follow, bold_prefix="追问时补充：")
    if boundary:
        add_callout("真实性边界", boundary, risk=True)


# Cover
for _ in range(3):
    doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("RGB-D人体动作识别与机械臂控制项目")
set_run(r, 25, bold=True, color=INK)
p.paragraph_format.space_after = Pt(8)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("机器人控制系统开发岗位｜面试细节问答")
set_run(r, 16, bold=True, color=BLUE)
p.paragraph_format.space_after = Pt(28)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("围绕 RealSense D435、MediaPipe、ROS2、StarAI双臂控制与示教数据")
set_run(r, 11.5, color=MUTED)
p.paragraph_format.space_after = Pt(58)
add_callout("使用方法", "先用文档中的项目口径形成稳定回答，再把所有【待确认】参数替换成代码或日志中的真实值。面试中先回答实际实现，再说明工程化改进，不要把改进方案说成已经完成。")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("适用岗位：机器人控制系统开发 / 机器人软件开发 / 机器人视觉系统集成")
set_run(r, 10, color=MUTED)
doc.add_page_break()

doc.add_heading("使用前必须统一的项目口径", level=1)
add_para("项目本质：硬件采用StarAI双臂系统，由Violin Leader与Cello Follower构成主从遥操作链路；在此基础上，使用D435和MediaPipe识别挥手、抬手、指向等离散人体动作，通过ROS2传递动作事件，由控制节点映射为Cello跟随臂的预定义关节动作。")
add_para("已完成重点：RGB-D采集与对齐、关键点提取、时序动作判断、ROS2通信闭环，以及基于厂商串口SDK的StarAI设备接入、关节读写封装、动作约束和示教数据组织。")
add_para("不能夸大的部分：当前项目不是把33个人体关键点直接实时重定向为机器人全部关节角；如果没有完成手眼标定、在线IK、力矩控制或模仿学习训练，需要明确说明属于后续方案。")

doc.add_heading("回答结构", level=2)
add_code("1. 先说输入、输出和坐标系\n2. 再说具体API、消息和数据流\n3. 然后说异常处理与安全约束\n4. 最后说验证方式和当前边界")

doc.add_heading("第一部分　RealSense D435与RGB-D处理", level=1)
add_qa("项目中是怎样采集RGB和深度图的？",
       "使用Intel RealSense D435同时获取彩色帧和深度帧。若直接使用SDK，流程是配置color和depth stream、启动pipeline、等待frameset并检查两帧是否有效；若通过ROS2驱动，则订阅彩色图、对齐深度图和CameraInfo。每一帧数据必须带帧号和采集时间戳。",
       "需要给出真实分辨率、帧率、图像编码和SDK/ROS2版本，例如640×480、30 FPS只是常见配置，不能当作未验证的项目事实。",
       "请在面试前从代码确认：使用pyrealsense2还是realsense2_camera；分辨率、FPS、序列号和深度单位。")
add_qa("为什么彩色图和深度图需要对齐？",
       "D435的彩色与深度来自不同传感器，两者的内参、外参、分辨率和视场角不同。相同像素下标并不对应同一条空间射线。项目把深度映射到彩色相机像素坐标系，使MediaPipe在彩色图上得到的(u,v)能够查询同一空间点的深度。",
       "几何过程是：深度像素用深度内参反投影到深度相机三维坐标，再用深度到彩色相机外参变换，最后用彩色内参投影到彩色图。")
add_qa("项目中具体如何完成对齐？",
       "直接使用RealSense SDK时，可通过rs.align(rs.stream.color)处理frameset，再分别读取aligned depth frame和color frame；通过ROS2驱动时，订阅aligned_depth_to_color图像。对齐后的深度图处在彩色图像的像素几何下。",
       "对齐后的深度查询使用彩色图坐标；将像素反投影为三维点时也应使用对齐目标相机对应的内参。",
       code="align = rs.align(rs.stream.color)\naligned = align.process(frames)\ncolor = aligned.get_color_frame()\ndepth = aligned.get_depth_frame()")
add_qa("对齐之后是否意味着RGB和深度完全同步？",
       "不完全等价。空间对齐解决的是像素对应关系，时间同步解决的是采集时刻对应关系。RealSense frameset通常已经进行了设备侧配对；如果通过两个ROS2 Topic接收，还应根据header.stamp使用精确或近似时间同步，不能仅依赖两个回调各自的最新数据。",
       "人体运动时，时间错配会使手腕像素查询到背景或旧位置的深度，因此空间对齐和时间同步缺一不可。")
add_qa("怎样获得深度值，深度单位是什么？",
       "Z16深度图保存的是设备深度单位，使用SDK时可以调用get_distance(u,v)直接获得米；也可以读取原始值后乘depth_scale转换为米。应在记录数据时明确保存单位，避免毫米和米混用。",
       "对深度设置合理范围，并过滤0值、NaN和超量程数据。",
       "必须确认代码中使用get_distance还是原始数组乘depth_scale，不能只说‘读取像素值’。")
add_qa("人体目标距离具体是怎样计算的？",
       "先选取稳定的人体区域，例如左右肩或左右髋形成的躯干中心，而不是随意读取检测框中心单像素。然后在该点附近取局部深度窗口，剔除0值和超范围点，取有效深度中位数作为人体到相机的Z向距离。",
       "如果需要欧氏距离，则先反投影得到(X,Y,Z)，再计算sqrt(X²+Y²+Z²)。面试时必须明确项目报告的是Z向距离还是欧氏距离。")
add_qa("为什么不直接读取一个像素的深度？",
       "深度图存在空洞、飞点、遮挡边缘混合和反光噪声。单像素容易读到0或背景。局部窗口中位数对少量离群值更鲁棒，同时需要限制窗口落在人体躯干区域内，避免把背景混入。",
       "可选改进包括深度时域滤波、空间滤波以及基于人体掩膜的区域统计，但只能把实际使用过的方法说成已实现。")
add_qa("如何把二维关键点转换为相机坐标系三维点？",
       "MediaPipe归一化坐标先转换为彩色图像像素u=xW、v=yH，再从对齐深度图获得Z，最后使用彩色相机内参反投影：X=(u-cx)Z/fx，Y=(v-cy)Z/fy。得到的点属于camera_color_optical_frame。",
       "光学坐标系通常是x向右、y向下、z向前，不能直接当成机器人base_link坐标。",
       code="u = x_norm * width\nv = y_norm * height\nX = (u - cx) * Z / fx\nY = (v - cy) * Z / fy")
add_qa("MediaPipe输出的z能否直接当作D435真实深度？",
       "不能。MediaPipe landmark.z是模型估计的相对深度量，不是D435测得的米制相机Z坐标。需要真实三维控制时，应当使用对齐深度图和相机内参完成反投影。")
add_qa("图像做了缩放、裁剪、翻转后怎样查询深度？",
       "必须把推理图上的关键点逆映射回原始彩色图坐标，再去查询对齐深度。缩放要除以缩放比例，裁剪要加回ROI偏移，水平翻转要执行u_original=W-1-u_flip。否则取到的深度来自错误位置。",
       "如果项目未做ROI或翻转，可以明确说输入保持与原彩色帧同一几何关系，因此无需额外逆变换。")
add_qa("多模态数据是如何组织的？",
       "以一帧或一个时间戳为索引，将彩色图、对齐深度图、相机内参、人体关键点、关键点置信度、动作标签和valid标志关联起来。不是简单把文件放在同一目录，而是保证同一条记录能追溯到同一采集时刻和坐标系。",
       "某一模态缺失时应标记该帧无效或记录缺失字段，不能把上一帧数据无期限补到当前帧。")
add_qa("如何处理相机丢帧、深度空洞和设备断开？",
       "单帧缺失时跳过当前融合计算并记录原因；短时关键点或深度缺失可以允许状态机保持候选状态，但连续多帧缺失必须使人体动作无效。设备断开时控制端进入安全状态，不继续执行依赖实时视觉的新动作；恢复后重新初始化数据流并清空旧队列。")

doc.add_heading("第二部分　MediaPipe、OpenCV与动作解析", level=1)
add_qa("项目具体使用了MediaPipe的哪个模块，输出什么？",
       "使用MediaPipe Pose提取人体姿态，输出33个人体关键点。每个关键点包含归一化x、y、相对z以及visibility等信息。OpenCV负责图像格式转换、必要的几何处理、可视化和调试。",
       "如果实际使用的是Holistic或不同model_complexity，必须按代码回答；还要确认static_image_mode、检测置信度和跟踪置信度设置。")
add_qa("为什么要把OpenCV的BGR图转换为RGB？",
       "OpenCV读取和RealSense常用输出是BGR通道顺序，而MediaPipe接口按RGB解释输入。如果不转换，颜色通道错位会降低模型输入一致性和检测稳定性。",
       code="rgb = cv2.cvtColor(bgr, cv2.COLOR_BGR2RGB)\nresult = pose.process(rgb)")
add_qa("33个关键点是否全部用于动作判断？",
       "不需要。项目针对具体动作选择相关关键点，例如肩、肘、腕和髋。减少无关关键点可以使规则更清晰，但仍可利用躯干尺度做归一化，降低人与相机距离变化对阈值的影响。")
add_qa("抬手动作如何定义？",
       "可使用手腕相对肩膀的垂直位置、肘部位置和持续时间共同判断。例如图像坐标y向下，手腕y小于肩膀y表示手腕处于肩上方；同时要求相关关键点置信度达标，并连续多帧成立。",
       "具体阈值、左右手规则和确认帧数必须与项目代码一致。单一关系只能作为候选条件，不能直接触发机械臂。")
add_qa("挥手动作如何与普通抬手区分？",
       "先要求手腕处于肩部上方或附近，再在时间窗口内分析手腕横向位置或相对肩宽归一化位置的往复变化。只有出现方向反转并达到一定幅度和持续时间，才判断为挥手；静止抬手不会满足往复特征。")
add_qa("指向动作如何判断？",
       "使用肩、肘、腕形成的方向和肘关节伸展程度。可以计算肩到腕方向向量，并根据肩-肘-腕夹角判断手臂是否接近伸直，再结合持续时间和置信度确认动作。",
       "如果项目只是区分‘指向’类别，不要声称已把方向向量转换成机械臂末端三维目标。")
add_qa("关节夹角是怎么计算的？",
       "以肘关节为顶点，定义向量a=肩-肘、b=腕-肘，通过点积计算夹角：theta=acos((a·b)/(|a||b|))。计算前要检查向量长度，点积比值要裁剪到[-1,1]，避免浮点误差导致acos非法。")
add_qa("为什么使用相对尺度而不是固定像素阈值？",
       "同一个人在不同距离下的像素尺寸不同。可以用肩宽或躯干长度归一化手腕位移，例如手腕横向位移除以肩宽，使动作阈值对拍摄距离更稳定。固定像素阈值仅适合相机位置和人体距离基本固定的受控场景。")
add_qa("时序窗口具体解决什么问题？",
       "单帧关键点存在抖动和偶发误检，而动作本身具有持续时间和方向变化。时序窗口可以进行连续帧确认、多数投票、位移趋势或方向反转判断，从而降低单帧误触发。窗口过短不稳定，过长会增加响应延迟。")
add_qa("窗口长度应该怎么选？",
       "按照实际帧率和动作持续时间选择，最好按时间而不只是固定帧数。例如相机30 FPS时5帧约167 ms。应通过录制数据回放，在误触发率与响应延迟之间调节窗口、确认帧数和冷却时间。",
       "面试前必须确认真实窗口长度、通过条件和冷却时间；如果没有正式消融实验，可以说通过多轮实机回放调试确定，不要编造准确率。")
add_qa("什么是异常帧，项目如何处理？",
       "异常帧包括未检测到人体、关键点置信度过低、坐标越界、帧间跳变过大、深度为0、时间戳不连续和处理超时。异常帧不参与动作特征更新；短时缺失可保留有限状态，连续缺失则清空窗口并返回IDLE。")
add_qa("是否可以一直使用上一帧关键点补偿缺失？",
       "不可以无限保持。上一帧仅可用于极短时的容错，并应携带有效期限。长期沿用会把过期人体姿态当作当前状态，可能导致机械臂误动作。超过缺失阈值后应发布valid=false并重置动作状态。")
add_qa("置信度判断具体怎么做？",
       "先对参与当前规则的关键点逐点检查visibility或相应可靠度，任一关键点低于阈值时不计算该动作。动作级置信度还应综合关键点可靠度、规则满足程度和时序一致性。关键点置信度不等于动作分类准确率。")
add_qa("关键点使用了什么滤波？",
       "项目核心可以回答为时序窗口、异常值拒绝和连续帧确认。如果代码只做了这些，就不要声称实现卡尔曼滤波。进一步优化可采用指数滑动平均或One Euro Filter，但滤波参数需要在平滑程度与响应延迟之间权衡。",
       boundary="把‘可采用卡尔曼/One Euro’明确说成改进方案，除非代码中确实已经实现并验证。")
add_qa("如何验证动作识别规则有效？",
       "使用在线实机测试和离线录像回放，分别统计正确触发、漏触发、误触发和触发延迟；还应覆盖不同距离、左右手、局部遮挡和短时关键点丢失。控制联调前先让系统只记录动作事件，不连接机械臂，以隔离感知问题。")

doc.add_heading("第三部分　ROS2系统与通信", level=1)
add_qa("为什么项目使用ROS2，而不是直接在一个Python程序中调用机械臂？",
       "ROS2将相机采集、感知、动作判断和机械臂控制解耦，支持独立开发、跨进程通信、QoS配置、状态反馈、Rosbag记录与回放。后续更换相机或机械臂时，只要保持消息接口，其他模块无需整体重写。")
add_qa("ROS2节点如何划分？",
       "逻辑上分为相机采集、人体感知/动作识别、控制决策和机械臂驱动。原型中感知与动作识别可以在同一节点，控制和设备接口单独封装；工程化时再根据实时性和故障隔离拆分。",
       boundary="不要把逻辑架构图说成实际部署节点数。面试时先说真实节点划分，再说明理想拆分。",
       code="Camera → Perception/Action → Robot Control → LeRobot/Serial\n                         ↑                │\n                         └──── Status ────┘")
add_qa("项目有哪些Topic和消息？",
       "相机侧包括彩色图、对齐深度图和CameraInfo；感知侧发布人体关键点和动作结果；控制侧接收动作事件并发布机械臂指令或调用设备接口，同时发布执行状态。动作消息至少应包含header、action_id、confidence、valid和event_id。",
       boundary="如果原型实际使用Int32或String，应说‘原型使用简单消息验证闭环，产品化会改为自定义强类型消息’，不要反向虚构。")
add_qa("Topic、Service和Action分别适合什么？",
       "图像、关键点和关节状态是连续异步数据，适合Topic；设备复位、参数查询等一次请求一次响应适合Service；机械臂动作持续一段时间且需要反馈、取消和结果，适合Action。",
       "如果原型用Topic传动作编号，可以说明Action是工程化升级方向。")
add_qa("图像与控制命令的QoS为什么不同？",
       "图像高频且关注最新状态，通常使用Best Effort和小队列，允许丢旧帧以控制延迟。动作事件、控制指令和故障状态需要Reliable，减少传输丢失。但Reliable只保证DDS传输层尽力可靠，不代表机械臂一定执行成功，仍需业务反馈和超时。")
add_qa("发布端和订阅端QoS不兼容会怎样？",
       "节点可能发现彼此但无法建立有效通信，表现为Topic存在却收不到消息。排查时应查看topic info、端点QoS以及可靠性和Durability设置，而不是只检查Topic名称。")
add_qa("RGB和深度两个Topic如何同步？",
       "使用message_filters的TimeSynchronizer或ApproximateTimeSynchronizer，根据header.stamp匹配同一时刻附近的RGB和对齐深度。允许时间差应结合帧率设置；若差值超过阈值则丢弃该组合，不能随意拼接两张最新图。")
add_qa("人体姿态和动作结果使用哪个时间戳？",
       "结果消息的header.stamp继承原始图像采集时间，表示该结果对应的观测时刻。节点可另行记录推理完成、消息发布和机械臂开始执行时间，用于计算感知延迟和端到端延迟。")
add_qa("为什么不能只在推理结束时调用now()作为结果时间？",
       "那样记录的是结果发布时间，而不是人体被观测的时间。推理负载变化时，两个时间的差值也变化，会破坏RGB-D同步、TF查询、消息过期判断和延迟分析。")
add_qa("图像处理速度低于相机帧率怎么办？",
       "不让图像无限排队。图像订阅使用小队列，回调仅更新最新帧，推理线程处理最新可用数据，必要时主动丢弃旧帧。机器人在线感知更关注低延迟，而不是保证每一帧都被处理。")
add_qa("如何避免MediaPipe推理阻塞控制消息？",
       "将感知计算与控制状态、故障和急停回调分开。可使用独立线程或不同Callback Group配合MultiThreadedExecutor。图像回调不要做阻塞串口读写，串口通信使用独立线程和有界队列。")
add_qa("为什么动作消息需要event_id？",
       "DDS重传、节点重连、状态机重复发布都可能使同一动作到达多次。event_id作为业务事件唯一编号，控制节点记录已经处理的编号，重复事件直接拒绝，从业务层避免机械臂重复执行。")
add_qa("如何判断一条控制消息是否过期？",
       "控制节点用当前ROS时间减去消息携带的图像采集时间。如果超过允许最大年龄，则丢弃，不让机械臂执行已经过时的人体动作。阈值应结合感知、通信和机械臂响应延迟确定。")
add_qa("怎样验证和排查ROS2通信链路？",
       "使用ros2 node list检查节点，ros2 topic info/echo/hz/delay检查端点、内容、频率和延迟，rqt_graph查看连接，Rosbag记录图像、动作、命令和状态以复现偶发问题。日志中使用event_id串联一次动作从识别到完成的全过程。",
       code="ros2 topic hz /human_action\nros2 topic delay /camera/color/image_raw\nros2 topic info /robot/status\nrqt_graph")
add_qa("如果感知节点崩溃或ROS2通信中断，机械臂怎么办？",
       "控制端不能继续使用旧的人体状态。通过心跳或指令更新时间设置看门狗，超时后停止接受新动作、减速停止或保持安全位置并进入FAULT。通信恢复后清空旧队列，需要重新确认或复位，不能自动执行积压指令。")
add_qa("ROS2日志和Rosbag分别解决什么问题？",
       "结构化日志记录事件、状态迁移、指令和错误原因；Rosbag保存消息级原始数据，支持离线重放同一感知输入。二者结合可以区分感知误判、消息延迟、控制拒绝和设备执行故障。")

doc.add_heading("第四部分　机械臂动作映射、通信与控制", level=1)
add_qa("项目使用的具体机械臂配置是什么？",
       "项目使用StarAI双臂平台，采用Violin主臂（Leader）和Cello跟随臂（Follower）组成主从遥操作系统。Violin侧提供操作者示教的关节状态，Cello侧接收对应关节目标并执行；项目又在这一平台上增加了视觉识别控制链路。",
       boundary="面试中要说清Violin是主臂、Cello是跟随臂，不要只说泛化的‘LeRobot机械臂’。双臂具体端口号和左右臂命名应以项目配置为准。")
add_qa("传统主从遥操作链路与视觉控制链路有什么区别？",
       "主从遥操作时，系统读取Violin主臂各关节状态，并把关节目标发送给Cello跟随臂，实现人工直接示教。视觉控制时不依赖操作者持续拖动主臂，而是由D435和MediaPipe识别离散人体动作，ROS2控制节点查询动作映射表，再向Cello发送预定义关节目标。两条链路共用机械臂设备接口，但上层指令来源不同。",
       boundary="视觉链路实现的是离散动作到预定义关节动作的映射，不等同于视觉连续遥操作或完整人体运动重定向。")
add_qa("人体动作怎样变成机械臂动作？",
       "项目采用离散映射：先把人体关键点序列识别为挥手、抬手或指向等动作类别，再由ROS2控制节点查询动作—关节目标映射表，生成Cello双臂对应的预定义关节动作。控制节点完成有效性、忙状态和限位检查后，通过厂商串口SDK的关节写接口下发；感知节点不直接操作舵机。")
add_qa("为什么不把33个人体关键点直接映射为机器人关节角？",
       "人体与机械臂的自由度、关节结构、尺度和运动范围不同，直接映射会产生不可达、跳变和碰撞风险。连续运动重定向需要坐标标定、骨架映射、IK、约束优化和轨迹规划，不属于当前离散动作闭环的核心实现。")
add_qa("预定义动作保存的是关节角还是末端位姿？",
       "应按真实项目回答。如果保存目标关节角，控制端直接在关节空间插值；如果保存末端位姿，则还需要在线IK生成关节目标。当前项目更安全的口径是预定义动作或目标关节姿态，由底层接口执行。",
       boundary="面试前确认每个动作文件或配置的真实数据结构，避免把关节目标和末端Pose混为一谈。")
add_qa("项目是否使用了逆运动学？",
       "如果核心动作是预定义关节动作，可以明确回答没有依赖在线IK。若后续升级为相机三维目标控制，则需要把目标转换到base_link，求解IK并进行多解选择、关节限位、奇异性和碰撞检查。",
       boundary="除非代码中确实有IK求解器调用，不要声称完成了在线逆运动学。")
add_qa("动作平滑具体怎么做？",
       "从机械臂当前关节状态出发，在给定动作时长内逐步插值到目标关节状态，而不是一次跳变；每一步检查位置范围，并限制相邻指令变化。更完善的方案使用三次或五次多项式满足起终点速度、加速度连续。",
       boundary="如果项目实际只使用简单线性插值或设备库内置平滑，应明确说清楚，不要把改进算法说成已实现。",
       code="q(t) = (1 - α) q_current + α q_target,   α ∈ [0, 1]")
add_qa("为什么线性插值仍可能不够？",
       "线性位置曲线在起止时刻可能产生速度突变，对机械臂造成冲击。三次或五次轨迹可以设置起终点速度甚至加速度为0，使轨迹更平滑；实际还需要结合控制周期、速度和加速度上限进行时间参数化。")
add_qa("机械臂控制的是位置、速度还是力矩？",
       "本项目主要属于上层位置或目标关节姿态控制，底层舵机/驱动器完成电机闭环。项目没有直接实现电流环或力矩环。ROS2上层负责动作决策、目标生成、状态管理和设备通信。")
add_qa("机械臂执行过程中又收到新动作怎么办？",
       "普通动作在EXECUTING状态下不直接抢占，可选择拒绝或只保留最新待执行事件；急停和安全停止具有最高优先级。当前动作完成并经过冷却或状态确认后，系统才重新接受普通动作。")
add_qa("动作状态机有哪些状态？",
       "感知侧可以包含IDLE、CONFIRMING、TRIGGERED和COOLDOWN；控制侧包含IDLE、EXECUTING、FAULT和RESET。两侧职责不同：感知侧确认动作，控制侧以机械臂真实反馈决定是否完成，不能由感知节点假设执行成功。",
       code="IDLE → CONFIRMING → TRIGGERED → EXECUTING → COOLDOWN → IDLE\n                         任何异常 ─────────→ FAULT → RESET")
add_qa("冷却机制为什么需要？",
       "人体保持同一姿态时，感知节点可能连续多帧都判断为同一动作。冷却机制使一次动作只触发一次，并要求动作退出或冷却时间结束后才能重新触发。event_id进一步防止消息层面的重复执行。")
add_qa("机械臂设备是如何接入的？",
       "项目使用厂商提供的串口SDK完成Violin主臂和Cello跟随臂的设备接入。在SDK关节读取、关节写入接口之上封装ROS2机械臂控制节点：节点接收上层动作事件，完成动作查表、单位与关节顺序转换、安全检查和指令下发，并发布机械臂状态。",
       boundary="底层总线舵机协议、电机闭环和SDK不是本人从零实现；个人工作重点是SDK接入、ROS2节点封装、动作映射、状态约束与联调。")
add_qa("主从遥操作具体如何实现？",
       "初始化时分别打开Violin Leader和Cello Follower对应的串口并加载校准配置。运行时周期性读取Violin各关节位置，按左右臂和关节编号进行映射、单位转换与限幅，再通过SDK写入Cello对应关节；同时读取Cello反馈用于监控跟随状态。",
       boundary="控制周期、端口名、SDK函数名、是否逐关节写入或批量同步写入，必须按实际代码回答。")
add_qa("为什么Violin和Cello都要做关节校准？",
       "主从臂的编码器零点、方向、量程和机械安装关系可能不同。校准用于确定各关节的有效范围和标准化映射，避免同一数值在主臂与跟随臂上对应不同姿态。校准后仍需核对左右臂顺序、关节方向和限位。",
       boundary="如果实际使用了厂商或LeRobot生成的校准文件，应说明调用并验证该机制，不要声称自行设计了完整标定算法。")
add_qa("双臂串口如何避免左右臂或主从臂混淆？",
       "启动时通过配置显式绑定role、side和port，例如leader_left、leader_right、follower_left、follower_right；连接后读取设备信息或执行低风险关节检查确认身份。控制层不依赖系统临时分配的tty序号推断角色，并在日志中记录端口和设备映射。",
       boundary="实际是否使用固定/dev/ttyUSB编号、udev规则或设备序列号，需要从部署配置确认。")
add_qa("串口数据包通常怎样组织？",
       "项目通过厂商串口SDK调用关节读写接口，应用层主要处理端口配置、设备ID、关节编号、目标值、超时和错误码。帧头、命令字、校验等底层报文由SDK封装；只有阅读过协议或抓包验证后，才进一步解释具体报文字段。",
       "通用串口协议可能包含帧头、长度、命令字、Payload和校验，但这不是可以直接冒充为项目实际字段的答案。",
       code="Header | Length | Command | Sequence | Payload | CRC")
add_qa("串口读写为什么不放在ROS2订阅回调里？",
       "阻塞式串口读写会占用Executor线程，使控制状态、故障甚至急停回调延迟。更合理的做法是订阅回调完成校验后把命令放入有界线程安全队列，独立通信线程负责编码、发送、超时和应答解析。")
add_qa("串口超时重传会不会让机械臂执行两次？",
       "可能。因此通信重传不能被设备解释为新的业务动作。可以使用相同sequence/event_id，使设备或控制层识别重复包；也可以先查询执行状态再决定是否重发。传输可靠性和业务幂等性需要分别设计。")
add_qa("如何判断机械臂动作执行完成？",
       "优先使用设备状态反馈，同时比较实际关节角与目标角误差，并确认关节速度接近0、条件连续多个周期成立。还要设置动作超时，超时后进入故障或安全停止，而不是无限等待。",
       boundary="如果原型只使用固定等待时间，需要承认这是验证阶段做法，并说明反馈式到位判断是改进方案。",
       code="max_i |q_measured[i] - q_target[i]| < εq\nmax_i |dq_measured[i]| < εv")
add_qa("项目中的状态约束具体是什么？",
       "只有IDLE接受普通动作；EXECUTING期间拒绝重复动作；FAULT禁止运动；控制消息必须未过期且event_id未处理；目标关节位置、指令步长、速度和动作时长必须在允许范围；通信或执行超时触发安全状态。")
add_qa("视觉误识别时如何降低机械臂风险？",
       "建立多层防线：关键点置信度过滤、连续帧确认、动作状态机、event_id去重、机械臂忙状态、关节/速度限制、执行超时和急停。感知结果只是候选控制意图，最终是否执行由控制节点在检查系统状态后决定。")
add_qa("如果需要让机械臂连续跟随人体手部三维位置，流程是什么？",
       "从对齐RGB-D得到手部在相机坐标系的三维点，通过相机到机械臂基座外参转换到base_link，生成末端目标，再进行工作空间检查、IK、轨迹规划、碰撞检查和控制执行。还需要滤波、限速和失联保护。",
       code="p_base = T_base_camera · p_camera")
add_qa("当前项目是否需要相机到机械臂基座的外参？",
       "离散人体动作映射到预定义机械臂动作时，不需要利用人体三维点生成机械臂空间目标，因此核心闭环不依赖相机-基座外参。只有升级为三维目标点或末端位姿控制时，才必须进行眼在手外标定。",
       boundary="不要把多相机标定经验等同于已经完成D435与机械臂的手眼标定。")

doc.add_heading("第五部分　示教数据、时间戳和动作复现", level=1)
add_qa("一条示教数据具体包含什么？",
       "以Episode为单位记录时间戳、RGB/深度观测、人体关键点、动作标签、下发关节目标、机械臂实际关节状态、执行状态、成功/失败标记和配置元数据。至少要能从一次Episode追溯到观测、指令和实际执行结果。")
add_qa("为什么既要记录目标关节角，也要记录实际关节角？",
       "目标关节角表示控制意图，实际关节角反映机械臂真实响应，两者受到跟踪误差、通信延迟和机械约束影响。只记录命令无法判断机器人是否实际到位，也不利于训练和复现。")
add_qa("为什么示教数据需要统一时间戳？",
       "模仿学习需要建立时刻t的观测与对应动作之间的关系。若图像、人体动作和机器人状态错位，模型会学到错误映射。统一时间轴支持重采样、延迟补偿、数据检查和回放复现。")
add_qa("图像30 Hz、关节状态更高频时怎样对齐？",
       "选定统一训练时间轴，对图像采用最近邻或最近有效帧，对连续关节状态进行时间插值；同时设置最大允许时间差，超过阈值的样本标记无效。不能无条件把不同频率数据按数组下标对应。")
add_qa("Observation和Action是否应该使用同一时刻？",
       "需要根据控制延迟定义。训练样本通常是观测o_t对应策略应输出的动作a_t，但真实系统中采集、推理和执行存在延迟，因此要记录各阶段时间戳，并明确数据集使用命令生成时刻、下发时刻还是执行时刻。")
add_qa("一个动作片段如何划分？",
       "可以以状态机事件为边界：候选动作确认前保留少量上下文，TRIGGERED或控制开始作为核心起点，机械臂到位或超时作为结束。成功与失败片段都可以保存，但必须通过状态字段区分。")
add_qa("运行日志记录哪些字段？",
       "记录frame_id、capture_stamp、action_id、confidence、event_id、控制命令、机械臂状态、串口发送/应答、动作开始结束、拒绝原因、超时和故障码，并保存软件和配置版本，使一次误触发可以完整追踪。")
add_qa("怎样利用日志复现一次错误？",
       "先用event_id定位动作事件，查看对应图像和关键点，再检查时序窗口、置信度和状态迁移，随后检查控制节点是否因过期、忙状态或限位拒绝，最后查看串口应答和机械臂反馈。Rosbag可重放感知输入，验证问题是否可重复。")
add_qa("这些数据是否已经用于训练模仿学习模型？",
       "如果尚未完成训练，应回答当前完成的是采集链路、Episode组织、时间戳、动作片段和日志规范，为后续训练提供数据基础；没有把数据基础描述成已经训练并部署了ACT或Diffusion Policy。",
       boundary="简历表述是‘为后续模仿学习训练提供数据基础’，面试时保持这一边界。")
add_qa("如何检查示教数据质量？",
       "检查图像是否可解码、时间戳是否单调、多模态时间差是否超阈值、关键点和深度有效率、关节角是否越界、命令与反馈误差、Episode是否完整以及成功/失败标签是否一致。发现异常数据应隔离而不是直接混入训练集。")

doc.add_heading("第六部分　主管压力题与综合回答", level=1)
add_qa("这个项目最难的地方是什么？",
       "最难的不是调用MediaPipe，而是把不稳定的视觉结果安全地变成机械臂动作。项目通过置信度过滤、时序窗口、异常帧处理、动作状态机、event_id、忙状态和冷却机制，降低关键点抖动和重复消息对控制的影响。")
add_qa("这个项目中ROS2真正解决了什么？",
       "它解决了模块边界、异步通信、QoS、时间戳、状态反馈和数据回放问题，使感知、控制与设备接口可以独立调试和替换。ROS2本身不替代动作识别和安全控制逻辑。")
add_qa("系统实时性怎么证明？",
       "应分别记录采集、推理完成、动作发布、控制接收和执行开始时间，统计平均、P95和最大延迟，同时观察帧队列是否积压。若项目没有形成正式统计，应坦诚只完成交互式在线运行和功能联调，后续会补充系统化延迟报告。",
       boundary="不要在没有日志依据时给出具体P95、FPS或‘实时’毫秒数。")
add_qa("如果动作误触发率高，你按什么顺序排查？",
       "先离线回放确认关键点是否稳定，再检查坐标逆映射、深度有效性和置信度；然后检查特征阈值、窗口长度、连续帧确认和动作间互斥；最后检查ROS2是否重复发布、event_id去重和控制状态机。按层排查比直接调机械臂更安全。")
add_qa("如果重新做一次，你会改进什么？",
       "把动作消息改为强类型自定义消息，机械臂任务改为ROS2 Action；分离感知与控制Callback Group；加入统一心跳、诊断、P95延迟统计和故障注入测试；控制关键节点逐步迁移到C++；预定义动作改为可版本化轨迹配置。")
add_qa("你在项目中真正负责了哪些部分？",
       "应按代码和任务记录明确回答：RGB-D采集与对齐、MediaPipe关键点处理、动作规则与状态机、ROS2消息链路，以及在厂商串口SDK关节读写接口上封装StarAI控制节点、完成Violin—Cello主从遥操作联调、视觉动作映射、状态约束、日志和示教数据组织。SDK内部协议、舵机驱动和电机闭环不能算成从零实现。")
add_qa("这个项目与机器人控制系统开发岗位有什么匹配？",
       "项目覆盖传感器输入、ROS2通信、状态机、上层控制、设备接口、异常处理和日志回放。优势是能够把视觉感知结果接入机器人执行链路；需要继续加强的方向是C++工程化、运动学、轨迹规划和底层实时控制。")
add_qa("用一分钟总结这个项目。",
       "项目硬件采用StarAI双臂平台，由Violin主臂和Cello跟随臂构成主从遥操作系统。我在厂商串口SDK的关节读写接口上封装ROS2控制节点，并增加D435与MediaPipe视觉控制链路：系统通过33个人体关键点、置信度、时序窗口和状态机识别离散动作，再将动作事件映射成Cello双臂的预定义关节动作。控制端加入事件去重、忙状态、冷却、关节限位和超时约束，同时记录图像、动作、指令和反馈，为回放与后续示教数据训练提供基础。")

doc.add_heading("附录A　面试前必须确认的真实参数", level=1)
add_callout("填写要求", "所有参数必须从代码、配置、日志或实机记录中确认。没有测量的数据写‘未正式统计’，不要临时编造。", risk=True)
params = [
    ("采集接口", "pyrealsense2 / realsense2_camera", "【待确认】"),
    ("RGB分辨率与FPS", "例如640×480 @ 30 FPS", "【待确认】"),
    ("深度分辨率与FPS", "必须与实际配置一致", "【待确认】"),
    ("深度单位", "get_distance返回米 / raw×depth_scale", "【待确认】"),
    ("深度统计区域", "关键点周围窗口或躯干ROI", "【待确认】"),
    ("MediaPipe模块", "Pose / Holistic", "Pose（建议核代码）"),
    ("model_complexity", "0 / 1 / 2", "【待确认】"),
    ("检测/跟踪置信度", "min_detection_confidence等", "【待确认】"),
    ("动作窗口长度", "帧数或毫秒", "【待确认】"),
    ("连续确认条件", "例如N帧中M帧", "【待确认】"),
    ("动作冷却时间", "秒", "【待确认】"),
    ("ROS2节点数", "按实际部署", "【待确认】"),
    ("动作消息类型", "Int32/String/自定义消息", "【待确认】"),
    ("QoS与队列深度", "图像/动作/控制分别说明", "【待确认】"),
    ("平均FPS与延迟", "平均/P95/最大值", "未测则说明"),
    ("机械臂型号与关节数", "主从臂分别说明", "StarAI双臂：Violin Leader＋Cello Follower；6+1自由度/臂"),
    ("主从遥操作接口", "Leader读取与Follower写入方式", "厂商串口SDK关节读写接口"),
    ("左右臂串口映射", "leader/follower × left/right", "【待确认】"),
    ("预定义动作格式", "关节目标/轨迹/末端Pose", "【待确认】"),
    ("平滑算法", "线性/三次/库内置", "【待确认】"),
    ("控制更新频率", "Hz", "【待确认】"),
    ("串口协议与波特率", "厂商SDK封装；确认波特率", "协议由厂商SDK封装，波特率【待确认】"),
    ("到位判断方式", "反馈误差/固定等待", "【待确认】"),
    ("Episode数量与时长", "如有统计", "【待确认】"),
]
table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"
headers = ["参数", "确认内容", "你的真实值"]
for i, text in enumerate(headers):
    cell = table.rows[0].cells[i]
    set_cell_shading(cell, LIGHT_BLUE)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_run(r, 10, bold=True, color=DARK_BLUE)
set_repeat_table_header(table.rows[0])
for row in params:
    cells = table.add_row().cells
    for i, text in enumerate(row):
        p = cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(text)
        set_run(r, 9.3, color=RISK_RED if "待确认" in text or "未测" in text else INK)
set_table_geometry(table, [2160, 4080, 3120])

doc.add_heading("附录B　面试回答自检", level=1)
checks = [
    "能否在90秒内画出Camera—Perception—Action—Control—Robot完整链路？",
    "每一个坐标值能否说明属于图像坐标、相机坐标还是base_link？",
    "每一个时间戳能否说明代表采集、发布还是执行时刻？",
    "能否区分实际使用的消息和理想的工程化消息？",
    "能否说明单帧误检、连续丢帧、消息过期和设备断开分别怎样处理？",
    "能否解释为什么Reliable不等于机械臂一定执行成功？",
    "能否明确当前项目是离散动作映射，而不是完整人体运动重定向？",
    "能否拿出真实窗口、阈值、冷却、帧率和通信参数？",
    "能否区分LeRobot已有能力与自己编写的模块？",
    "能否清楚区分Violin—Cello主从遥操作链路与视觉动作控制链路？",
    "能否说明下一步如何迁移到C++、Action、IK和连续轨迹控制？",
]
for idx, text in enumerate(checks, 1):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_run(r)

doc.add_heading("附录C　推荐开场口径", level=1)
add_callout("一分钟项目介绍", "项目使用StarAI双臂平台，由Violin主臂和Cello跟随臂组成主从遥操作系统。我在厂商串口SDK的关节读写接口上封装ROS2控制节点，并在此基础上增加RealSense D435与MediaPipe视觉控制链路：相机侧完成彩色与深度对齐，感知侧提取33个人体关键点，通过置信度、时序窗口和状态机识别挥手、抬手、指向等离散动作；动作事件经ROS2发送到控制节点，映射为Cello双臂的预定义关节动作。控制端加入事件去重、忙状态、冷却、关节限位和超时约束，并统一记录观测、指令和反馈时间戳，为系统回放和后续示教数据训练提供基础。")

# Avoid isolated headings and set core properties.
for p in doc.paragraphs:
    if p.style.name.startswith("Heading") or p.style.name == "Question":
        p.paragraph_format.keep_with_next = True
    p.paragraph_format.widow_control = True

doc.core_properties.title = "RGB-D人体动作识别与机械臂控制项目面试细节问答"
doc.core_properties.subject = "机器人控制系统开发岗位面试准备"
doc.core_properties.author = ""
doc.core_properties.keywords = "ROS2, RealSense D435, MediaPipe, StarAI, Violin, Cello, 双臂遥操作"

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(f"created {OUT} with {q_no} Q&A items")
